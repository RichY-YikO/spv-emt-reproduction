from pathlib import Path
import argparse
import ast
import hashlib
import io
import json
import subprocess
import sys
import tokenize
import csv
import math

ROOT=Path(__file__).resolve().parent
CODE=ROOT/'code'

def release_files():
    ignored={'outputs','__pycache__','.git','.venv','.idea','.vscode','JGR_Atmospheres_vector_figures'}
    generated={'CHECKSUMS.sha256','Pearson_validation.json','SI_artist_validation.json','replot_manifest.json'}
    for p in sorted(ROOT.rglob('*')):
        if not p.is_file():continue
        rel=p.relative_to(ROOT)
        if any(part in ignored for part in rel.parts) or p.name in generated or p.suffix in ['.pyc','.zip']:continue
        if rel.parts[:2]==('data','raw'):continue
        if rel.parts[0]=='code' and p.suffix in ['.pdf','.png']:continue
        if rel.parts[0]=='code' and p.name.endswith('.layout.json'):continue
        yield p

def write_checksums():
    lines=[hashlib.sha256(p.read_bytes()).hexdigest()+'  '+p.relative_to(ROOT).as_posix() for p in release_files()]
    (ROOT/'CHECKSUMS.sha256').write_text('\n'.join(lines)+'\n',encoding='utf-8')

def check_checksums():
    path=ROOT/'CHECKSUMS.sha256'
    if not path.exists():raise FileNotFoundError('Missing release checksums; use --refresh-checksums only after reviewing intentional changes.')
    rows=path.read_text(encoding='utf-8').splitlines()
    for row in rows:
        h,rel=row.split('  ',1)
        assert hashlib.sha256((ROOT/rel).read_bytes()).hexdigest()==h,rel
    return len(rows)

def check_sources():
    files=list(CODE.glob('*.py'))+[ROOT/'reproduce.py',ROOT/'verify_release.py']
    for p in files:
        text=p.read_text(encoding='utf-8')
        tree=ast.parse(text,filename=str(p))
        compile(tree,str(p),'exec')
        assert not any(t.type==tokenize.COMMENT for t in tokenize.generate_tokens(io.StringIO(text).readline)),str(p)
        for n in ast.walk(tree):
            if isinstance(n,(ast.Module,ast.FunctionDef,ast.AsyncFunctionDef,ast.ClassDef)):
                assert ast.get_docstring(n) is None,str(p)
        for n in ast.walk(tree):
            if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id=='Path' and n.args and isinstance(n.args[0],ast.Constant):
                value=n.args[0].value
                assert not (isinstance(value,str) and len(value)>2 and value[1]==':'),str(p)
    return len(files)

def check_files():
    manifest=json.loads((ROOT/'validation'/'source_data_sha256.json').read_text(encoding='utf-8'))
    for x in manifest:
        p=ROOT/x['file']
        assert p.is_file(),str(p)
        assert hashlib.sha256(p.read_bytes()).hexdigest()==x['sha256'],str(p)
    assert len(list((ROOT/'figures'/'pdf').glob('Fig*.pdf')))==14
    assert len(list((ROOT/'figures'/'png').glob('Fig*.png')))==14
    expected={'S_final_pearson.npz','T_basis_and_correlations.npz','S_top100.csv','T_top100.csv','table_S1.csv'}
    assert {p.name for p in (CODE/'fig1_fig2_results').iterdir()}==expected
    assert not (CODE/'Fig1a_no_orientation_20261014').exists()
    assert not list((ROOT/'figures').rglob('*.json'))
    from openpyxl import load_workbook
    with (CODE/'fig1_fig2_results'/'table_S1.csv').open(encoding='utf-8-sig',newline='') as stream:
        rows=list(csv.reader(stream))
    workbook=load_workbook(ROOT/'figures'/'TableS1.xlsx',read_only=True,data_only=True)
    try:
        actual=list(workbook.worksheets[0].iter_rows(values_only=True))
        assert tuple(rows[0])==actual[0]
        assert len(actual)==7
        for source,target in zip(rows[1:],actual[1:]):
            assert source[0]==target[0]
            assert all(math.isclose(float(a),float(b),rel_tol=0,abs_tol=1e-10) for a,b in zip(source[1:],target[1:]))
    finally:
        workbook.close()
    return len(manifest)

def main():
    parser=argparse.ArgumentParser(description='Verify release integrity, comment removal, scientific data, SI artists or vector PDFs.')
    parser.add_argument('--mode',choices=['files','numerical','si','pdf','all'],default='files')
    parser.add_argument('--pdf-dir',type=Path,default=ROOT/'figures'/'pdf')
    parser.add_argument('--refresh-checksums',action='store_true',help='Regenerate the release manifest after reviewing intentional metadata or source changes.')
    parser.add_argument('--layout-dir',type=Path,help='Override the PDF layout validation directory.')
    args=parser.parse_args()
    args.pdf_dir=args.pdf_dir.resolve()
    if args.layout_dir is not None:args.layout_dir=args.layout_dir.resolve()
    if args.refresh_checksums:write_checksums()
    print(json.dumps({'python_files_checked':check_sources(),'source_files_hash_verified':check_files(),'release_files_hash_verified':check_checksums()}),flush=True)
    tasks=[]
    if args.mode in ['numerical','all']:tasks.append([str(CODE/'verify_final_pearson.py')])
    if args.mode in ['si','all']:tasks.append([str(CODE/'verify_supplement_figures.py')])
    if args.mode in ['pdf','all']:tasks.append([str(CODE/'verify_publication.py'),'--output-root',str(args.pdf_dir)])
    if args.layout_dir is not None:
        for task in tasks:
            if Path(task[0]).name=='verify_publication.py':task.extend(['--layout-dir',str(args.layout_dir)])
    for task in tasks:subprocess.run([sys.executable,*task],check=True,cwd=ROOT)

if __name__=='__main__':
    main()
